"""Loan Register Cleaner — Flask backend.

Run from the project directory:
    source venv/bin/activate
    python app.py
Then open http://127.0.0.1:5000
"""
import io
import re
import statistics
import difflib
from datetime import date, datetime

import openpyxl
from flask import Flask, request, jsonify, send_file
from openpyxl.styles import Font, PatternFill

app = Flask(__name__, static_folder="static", static_url_path="")

# Adapt these aliases and normalisation rules for another branch/vendor.
FIELDS = ["branch", "customer_name", "loan_amount", "loan_date", "phone", "status"]
COLUMN_MAP = {
    "branch": ["branch", "branch_name", "branch name"],
    "customer_name": ["customer name", "customer_name", "name", "customer"],
    "loan_amount": ["loan amt", "loan amount", "loan_amount", "amount"],
    "loan_date": ["loan date", "loan_date", "date"],
    "phone": ["phone no", "phone_no", "phone", "mobile", "mobile_number", "mobile number"],
    "status": ["status", "loan status", "loan_status"],
    "notes": ["notes", "remarks", "note", "register note"],
}
STATUS = {
    "active": "Active", "live": "Active", "open": "Active",
    "closed": "Closed", "close": "Closed", "settled": "Closed",
    "npa": "NPA", "non performing": "NPA",
}
BRANCH_ALIASES = {}  # e.g. {"andheri w": "Andheri"}
BRANCH_NOISE = {"br", "br.", "branch"}
OUTLIER_DAYS = 180
MULT = {None: 1, "k": 1e3, "l": 1e5, "lac": 1e5, "lakh": 1e5,
        "lakhs": 1e5, "cr": 1e7, "crore": 1e7}
S = {"rows": [], "log": [], "file": "", "order": "DMY"}


def normalize_header(header):
    """Normalize headers so spaces, underscores, hyphens and case don't matter."""
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def issue(field, level, msg, suggest=None, kind=None):
    return {"field": field, "level": level, "msg": msg,
            "suggest": suggest, "kind": kind, "open": True}


def p_amount(raw):
    if raw is None or str(raw).strip() == "":
        return None, "Loan amount is missing in the register.", None
    if isinstance(raw, (int, float)):
        value, why = float(raw), None
    else:
        text = re.sub(r"₹|rs\.?|inr|,", "", str(raw).lower()).strip()
        match = re.fullmatch(r"(-?\d+(?:\.\d+)?)\s*(k|l|lac|lakhs?|cr|crore)?", text)
        if not match:
            return None, f'"{raw}" is not a number.', None
        value = float(match[1]) * MULT[match[2]]
        why = f'Expanded "{match[2]}" shorthand' if match[2] else "Removed currency symbol / commas"
    if value <= 0:
        return None, f"Amount is {value:,.0f}; a loan cannot be zero or negative.", None
    return (int(value) if value.is_integer() else value), None, why


def p_date(raw, order="DMY"):
    if raw is None or str(raw).strip() == "":
        return None, "Loan date is missing.", None
    if isinstance(raw, datetime):
        return raw.date().isoformat(), None, None
    if isinstance(raw, date):
        return raw.isoformat(), None, None
    text, why = str(raw).strip(), "Standardised date format"
    match = re.fullmatch(r"(\d{4})[-./](\d{1,2})[-./](\d{1,2})", text)
    if match:
        y, m, d = map(int, match.groups())
    else:
        match = re.fullmatch(r"(\d{1,2})[-./](\d{1,2})[-./](\d{2}|\d{4})", text)
        if match:
            a, b, y = map(int, match.groups())
            d, m = (a, b) if order == "DMY" else (b, a)
            if y < 100:
                y += 2000
                why = "Two-digit year read as 20xx"
            if a <= 12 and b <= 12 and a != b:
                why = f'Ambiguous date read as {"day/month" if order == "DMY" else "month/day"} (the rest of the file uses it)'
        else:
            parsed = None
            for fmt in ("%B %d, %Y", "%b %d, %Y", "%d %B %Y", "%d %b %Y"):
                try:
                    parsed = datetime.strptime(text, fmt)
                    break
                except ValueError:
                    continue
            if parsed is None:
                return None, f'Could not understand the date "{text}".', None
            y, m, d = parsed.year, parsed.month, parsed.day
    try:
        return date(y, m, d).isoformat(), None, why
    except ValueError:
        return None, f'"{text}" is not a real calendar date (day {d}, month {m}, year {y}).', None


def p_phone(raw):
    if raw is None or str(raw).strip() == "":
        return None, "Phone number is missing.", None
    text = str(int(raw)) if isinstance(raw, float) and raw.is_integer() else str(raw).strip()
    digits = re.sub(r"[\s\-().]", "", text)
    if digits.startswith("+91"):
        digits = digits[3:]
    elif len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    if not re.fullmatch(r"[6-9]\d{9}", digits):
        return None, f'"{text}" is not a valid 10-digit mobile number.', None
    return digits, None, None if digits == text else "Removed spaces, dashes or country code"


def p_status(raw):
    value = STATUS.get(re.sub(r"\s+", " ", str(raw or "")).strip().lower())
    if not value:
        return None, f'Status "{raw or ""}" is not Active / Closed / NPA.', None
    return value, None, "Standardised capitalisation" if value != raw else None


def p_name(raw):
    text = re.sub(r"\s+", " ", str(raw or "")).strip()
    if not text:
        return None, "Customer name is missing in the register.", None
    if text.isupper() or text.islower():
        return text.title(), None, "Fixed capitalisation"
    return text, None, "Removed extra spaces" if text != raw else None


def p_branch(raw):
    words = [w for w in re.sub(r"\s+", " ", str(raw or "")).strip().split()
             if w.lower() not in BRANCH_NOISE]
    text = " ".join(words).title()
    if not text:
        return None, "Branch is missing.", None
    text = BRANCH_ALIASES.get(text.lower(), text)
    return text, None, "Standardised branch name" if text != raw else None


PARSE = {"branch": p_branch, "customer_name": p_name, "loan_amount": p_amount,
         "phone": p_phone, "status": p_status}


def read_sheet(stream):
    workbook = openpyxl.load_workbook(stream, data_only=True)
    worksheet = workbook.worksheets[0]
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The Excel workbook is empty.")
    headers = [str(value or "").strip() for value in rows[0]]
    normalized = [normalize_header(value) for value in headers]
    indexes = {}
    for field, aliases in COLUMN_MAP.items():
        alias_keys = {normalize_header(alias) for alias in aliases}
        indexes[field] = next(
            (i for i, header in enumerate(normalized) if header and header in alias_keys),
            None,
        )
    missing = [field for field in FIELDS if indexes.get(field) is None]
    if missing:
        found = ", ".join(h for h in headers if h)
        raise ValueError(
            "Could not find these columns: " + ", ".join(missing)
            + (f". Headers found: {found}" if found else ". No headers found.")
        )
    output = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(value not in (None, "") for value in row):
            continue
        output.append((row_number, {
            field: row[index] if index is not None and index < len(row) else None
            for field, index in indexes.items()
        }))
    return output


def detect_order(raw_rows, pick):
    if pick in ("DMY", "MDY"):
        return pick
    dmy = mdy = 0
    for _, row in raw_rows:
        match = re.fullmatch(r"(\d{1,2})[-./](\d{1,2})[-./]\d{2,4}",
                             str(row.get("loan_date") or "").strip())
        if match:
            dmy += int(match[1]) > 12
            mdy += int(match[2]) > 12
    return "MDY" if mdy > dmy else "DMY"


def same_person(a, b):
    x, y = a.lower().replace(".", ""), b.lower().replace(".", "")
    if difflib.SequenceMatcher(None, x, y).ratio() >= .85:
        return True
    first, second = x.split(), y.split()
    return (len(first) > 1 and len(second) > 1 and first[-1] == second[-1]
            and first[0][0] == second[0][0]
            and 1 in (len(first[0]), len(second[0])))


def process(raw_rows, order):
    processed = []
    for source_row, raw in raw_rows:
        row = {"id": len(processed), "src": source_row,
               "raw": {key: "" if value is None else str(value) for key, value in raw.items()},
               "clean": {}, "issues": [], "excluded": False}
        for field in FIELDS:
            parser = PARSE.get(field) or (lambda value: p_date(value, order))
            value, problem, reason = parser(raw.get(field))
            row["clean"][field] = value
            if problem:
                row["issues"].append(issue(field, "error", problem))
            elif reason:
                S["log"].append({"row": source_row, "field": field,
                                 "before": row["raw"][field], "after": value,
                                 "why": reason, "by": "Auto"})
        try:
            amount = float(str(raw.get("loan_amount")).replace(",", "").replace("₹", ""))
            if amount < 0:
                for item in row["issues"]:
                    if item["field"] == "loan_amount":
                        item["suggest"] = abs(amount)
        except (ValueError, TypeError):
            pass
        if re.match(r"^[A-Za-z]\.?\s", row["clean"]["customer_name"] or ""):
            row["issues"].append(issue("customer_name", "warn",
                                       "Only an initial is written for the first name; identity is uncertain."))
        processed.append(row)
    cross_checks(processed)
    return processed


def cross_checks(rows):
    names = {row["clean"]["branch"] for row in rows if row["clean"]["branch"]}
    for row in rows:
        branch = row["clean"]["branch"]
        base = next((name for name in names if branch and branch != name
                     and branch.startswith(name + " ")), None)
        if base:
            row["issues"].append(issue("branch", "warn",
                f'"{branch}" may be a different branch/area from "{base}".', base))

    dates = [(row, date.fromisoformat(row["clean"]["loan_date"]))
             for row in rows if row["clean"]["loan_date"]]
    if dates:
        median = statistics.median(value.toordinal() for _, value in dates)
        for row, value in dates:
            if value > date.today():
                row["issues"].append(issue("loan_date", "error",
                                           f"Loan date {value} is in the future."))
            elif abs(value.toordinal() - median) > OUTLIER_DAYS:
                typical = date.fromordinal(int(median))
                row["issues"].append(issue("loan_date", "warn",
                    f"{value} is far from the other loans (typically around {typical}). Check for a typo."))

    for row in rows:
        phone = row["clean"]["phone"]
        if phone and (len(set(phone)) <= 2 or phone.endswith("00000")
                      or phone in ("9876543210", "1234567890")):
            row["issues"].append(issue("phone", "warn",
                                       f"{phone} looks like a placeholder / dummy number."))

    for index, first in enumerate(rows):
        for second in rows[index + 1:]:
            a, b = first["clean"], second["clean"]
            if not a["customer_name"] or not b["customer_name"]:
                continue
            if not same_person(a["customer_name"], b["customer_name"]):
                continue
            same_loan = (a["loan_amount"] is not None
                         and a["loan_amount"] == b["loan_amount"]
                         and a["loan_date"] == b["loan_date"])
            same_phone = a["phone"] and a["phone"] == b["phone"]
            if not (same_loan or same_phone):
                continue
            details = []
            details.append("name" if a["customer_name"] == b["customer_name"] else "similar name")
            if same_loan:
                details.extend(["amount", "date"])
            if same_phone:
                details.append("phone")
            if same_loan:
                message = (f"Looks like a duplicate of {b['customer_name']} "
                           f"(row {second['src']}): matching " + ", ".join(details) + ".")
                kind = "dup"
            else:
                message = (f"Same customer (name + phone) as {b['customer_name']} "
                           f"(row {second['src']}) but a different loan. A second loan, or a duplicate?")
                kind = "repeat"
            first["issues"].append(issue("row", "warn", message, kind=kind))
            second["issues"].append(issue("row", "warn",
                message.replace(b["customer_name"], a["customer_name"])
                       .replace(str(second["src"]), str(first["src"])), kind=kind))

    by_phone = {}
    for row in rows:
        phone = row["clean"]["phone"]
        if phone:
            by_phone.setdefault(phone, []).append(row)
    for phone, group in by_phone.items():
        for row in group:
            others = [other for other in group if other is not row
                      and not same_person(other["clean"]["customer_name"] or "",
                                          row["clean"]["customer_name"] or "#")]
            if others:
                row["issues"].append(issue("phone", "warn",
                    "Same phone number is also on: " + ", ".join(
                        f"{other['clean']['customer_name'] or '(no name)'} (row {other['src']})"
                        for other in others) + ". Family member, or a wrong number?"))


def row_state(row):
    if row["excluded"]:
        return "left_out"
    return "review" if any(item["open"] for item in row["issues"]) else "ready"


def payload():
    return {"file": S["file"], "order": S["order"],
            "rows": [dict(row, state=row_state(row)) for row in S["rows"]],
            "changes": len(S["log"])}


@app.route("/")
def home():
    return app.send_static_file("index.html")


@app.post("/api/upload")
def upload():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify(error="Choose an Excel file to upload."), 400
    if not file.filename.lower().endswith(".xlsx"):
        return jsonify(error="Please upload an .xlsx Excel file."), 400
    try:
        raw_rows = read_sheet(file.stream)
        S.update(rows=[], log=[], file=file.filename,
                 order=detect_order(raw_rows, request.form.get("order", "auto")))
        S["rows"] = process(raw_rows, S["order"])
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else (
            "This does not look like a valid Excel (.xlsx) register."
        )
        return jsonify(error=message), 400
    return jsonify(payload())


def get_row(index):
    return S["rows"][int(index)]


@app.post("/api/fix")
def fix():
    data = request.get_json(force=True)
    try:
        row = get_row(data["id"])
        field = data["field"]
        if field not in FIELDS:
            return jsonify(error="Unknown field."), 400
        parser = PARSE.get(field) or (lambda value: p_date(value, S["order"]))
        value, problem, _ = parser(data.get("value"))
        if problem:
            return jsonify(error=problem), 400
        row["clean"][field] = value
        S["log"].append({"row": row["src"], "field": field,
                         "before": row["raw"].get(field, ""),
                         "after": value, "why": "Corrected by staff", "by": "Staff"})
        for item in row["issues"]:
            if item["field"] == field:
                item["open"] = False
        return jsonify(payload())
    except (KeyError, ValueError, IndexError, TypeError):
        return jsonify(error="Could not apply that correction. Please refresh and try again."), 400


@app.post("/api/approve")
def approve():
    data = request.get_json(force=True)
    try:
        row = get_row(data["id"])
        item = row["issues"][int(data["issue"])]
        if item["level"] == "error":
            return jsonify(error="This needs a corrected value, not an approval."), 400
        item["open"] = False
        S["log"].append({"row": row["src"], "field": item["field"],
                         "before": "", "after": "",
                         "why": "Staff confirmed: " + item["msg"], "by": "Staff"})
        return jsonify(payload())
    except (KeyError, ValueError, IndexError, TypeError):
        return jsonify(error="Could not approve this item. Please refresh and try again."), 400


@app.post("/api/exclude")
def exclude():
    data = request.get_json(force=True)
    try:
        row = get_row(data["id"])
        row["excluded"] = bool(data["value"])
        S["log"].append({"row": row["src"], "field": "row",
                         "before": "", "after": "",
                         "why": "Left out of clean file by staff" if row["excluded"]
                                else "Put back by staff", "by": "Staff"})
        return jsonify(payload())
    except (KeyError, ValueError, IndexError, TypeError):
        return jsonify(error="Could not update this row. Please refresh and try again."), 400


@app.get("/api/export")
def export():
    workbook = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    def add_sheet(sheet, headers, data):
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font, cell.fill = header_font, header_fill
        for values in data:
            sheet.append(values)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(
                60, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    ready = [row for row in S["rows"] if row_state(row) == "ready"]
    add_sheet(workbook.active, FIELDS,
              [[row["clean"][field] for field in FIELDS] for row in ready])
    workbook.active.title = "Clean Data"
    for row_number in range(2, len(ready) + 2):
        workbook.active.cell(row_number, 5).number_format = "@"

    review = [row for row in S["rows"] if row_state(row) != "ready"]
    add_sheet(workbook.create_sheet("Needs Review"),
              ["source_row", "state"] + FIELDS + ["why_flagged", "register_note"],
              [[row["src"], "Left out by staff" if row["excluded"] else "Needs review"]
               + [row["clean"][field] for field in FIELDS]
               + [" | ".join(item["msg"] for item in row["issues"] if item["open"]),
                  row["raw"].get("notes", "")]
               for row in review])
    add_sheet(workbook.create_sheet("Change Log"),
              ["source_row", "field", "before", "after", "reason", "by"],
              [[entry["row"], entry["field"], entry["before"], entry["after"],
                entry["why"], entry["by"]] for entry in S["log"]])

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = "cleaned_" + (S["file"] or "register.xlsx")
    return send_file(buffer, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)
